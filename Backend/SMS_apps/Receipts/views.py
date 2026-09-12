from django.db.models import Sum, Count
from django.http import HttpResponse

from rest_framework import viewsets, status
from rest_framework.decorators import action
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .models import (
    Receipt,
    ReceiptItem,
    Department,
    SubDepartment,
)

from .serializer import (
    ReceiptSerializer,
    ReceiptItemSerializer,
    DepartmentSerializer,
    SubDepartmentSerializer,
    DepartmentSummarySerializer,
)


# ---------------------------------------------------------------------------
# Shared helper: builds the styled .xlsx header block used by both exports.
# ---------------------------------------------------------------------------

def _styled_workbook(
    sheet_title,
    subtitle,
    headers,
    column_widths,
):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_title

    worksheet["A1"] = "School I/O System"
    worksheet["A2"] = subtitle

    worksheet["A1"].font = Font(
        bold=True,
        size=16,
    )

    worksheet["A2"].font = Font(
        bold=True,
        size=13,
    )

    header_row = 4

    for column, header in enumerate(
        headers,
        start=1,
    ):
        cell = worksheet.cell(
            row=header_row,
            column=column,
            value=header,
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="343A40",
        )

        cell.alignment = Alignment(
            horizontal="center",
        )

    for index, width in enumerate(
        column_widths,
        start=1,
    ):
        worksheet.column_dimensions[
            get_column_letter(index)
        ].width = width

    worksheet.freeze_panes = "A5"

    return (
        workbook,
        worksheet,
        header_row,
    )


# ===========================================================================
# DEPARTMENT VIEWSET
# ===========================================================================

class DepartmentViewSet(viewsets.ModelViewSet):
    """
    Full CRUD for Departments.

    Endpoints:

        GET    /api/departments/
        POST   /api/departments/
        GET    /api/departments/<id>/
        PUT    /api/departments/<id>/
        PATCH  /api/departments/<id>/
        DELETE /api/departments/<id>/

    Additional endpoint:

        GET /api/departments/<id>/subdepartments/

    Departments are database-managed rather than hard-coded choices.
    """

    queryset = (
        Department.objects
        .prefetch_related("subdepartments")
        .all()
    )

    serializer_class = DepartmentSerializer

    permission_classes = [
        IsAuthenticated,
    ]

    # ------------------------------------------------------------------
    # QUERYSET
    # ------------------------------------------------------------------

    def get_queryset(self):
        """
        Return active departments by default.

        To include inactive departments:

            /api/departments/?include_inactive=true
        """

        queryset = super().get_queryset()

        include_inactive = (
            self.request.query_params.get(
                "include_inactive"
            )
        )

        if include_inactive != "true":
            queryset = queryset.filter(
                is_active=True
            )

        return queryset.order_by("name")

    # ------------------------------------------------------------------
    # SUBDEPARTMENTS
    # ------------------------------------------------------------------

    @action(
        detail=True,
        methods=["get"],
        url_path="subdepartments",
    )
    def subdepartments(
        self,
        request,
        pk=None,
    ):
        """
        Return only the active subdepartments belonging
        to the selected department.

        Example:

            GET /api/departments/5/subdepartments/
        """

        department = self.get_object()

        subdepartments = (
            department.subdepartments
            .filter(is_active=True)
            .order_by("name")
        )

        serializer = SubDepartmentSerializer(
            subdepartments,
            many=True,
        )

        return Response(
            serializer.data
        )


# ===========================================================================
# SUBDEPARTMENT VIEWSET
# ===========================================================================

class SubDepartmentViewSet(viewsets.ModelViewSet):
    """
    Full CRUD for SubDepartments.

    Endpoints:

        GET    /api/subdepartments/
        POST   /api/subdepartments/
        GET    /api/subdepartments/<id>/
        PUT    /api/subdepartments/<id>/
        PATCH  /api/subdepartments/<id>/
        DELETE /api/subdepartments/<id>/

    Filtering:

        GET /api/subdepartments/?department=5
        GET /api/subdepartments/?is_active=true
    """

    queryset = (
        SubDepartment.objects
        .select_related("department")
        .all()
    )

    serializer_class = SubDepartmentSerializer

    permission_classes = [
        IsAuthenticated,
    ]

    # ------------------------------------------------------------------
    # QUERYSET / FILTERING
    # ------------------------------------------------------------------

    def get_queryset(self):
        """
        Filter subdepartments by department and/or active status.
        """

        queryset = super().get_queryset()

        department = (
            self.request.query_params.get(
                "department"
            )
        )

        is_active = (
            self.request.query_params.get(
                "is_active"
            )
        )

        if department:
            queryset = queryset.filter(
                department_id=department
            )

        if is_active == "true":
            queryset = queryset.filter(
                is_active=True
            )

        elif is_active == "false":
            queryset = queryset.filter(
                is_active=False
            )

        return queryset.order_by(
            "department__name",
            "name",
        )


# ===========================================================================
# RECEIPT VIEWSET
# ===========================================================================

class ReceiptViewSet(viewsets.ModelViewSet):
    """
    Full CRUD for Receipt.

    Receipt lookup uses receipt_id rather than the numeric database ID.

    Example:

        /api/receipts/RP-2026-00001/
    """

    queryset = (
        Receipt.objects
        .select_related("recorded_by")
        .prefetch_related("items")
        .all()
    )

    serializer_class = ReceiptSerializer

    permission_classes = [
        IsAuthenticated,
    ]

    lookup_field = "receipt_id"

    # ------------------------------------------------------------------
    # CREATE RECEIPT
    # ------------------------------------------------------------------

    def perform_create(self, serializer):
        """
        Always set recorded_by from the authenticated user.

        The client cannot choose another user.
        """

        serializer.save(
            recorded_by=self.request.user
        )

    # ------------------------------------------------------------------
    # LIST RECEIPTS
    # ------------------------------------------------------------------

    def list(
        self,
        request,
        *args,
        **kwargs,
    ):
        """
        Return receipts together with total expenditure.

        Receipt.total is maintained automatically by ReceiptItem,
        so summing Receipt.total is safe here.
        """

        queryset = self.filter_queryset(
            self.get_queryset()
        )

        total_expenditure = (
            queryset.aggregate(
                total=Sum("total")
            )["total"]
            or 0
        )

        page = self.paginate_queryset(
            queryset
        )

        serializer = self.get_serializer(
            page if page is not None else queryset,
            many=True,
        )

        data = serializer.data

        if page is not None:
            paginated = self.get_paginated_response(
                data
            )

            paginated.data[
                "total_expenditure"
            ] = total_expenditure

            return paginated

        return Response(
            {
                "results": data,
                "total_expenditure": total_expenditure,
            }
        )

    # ------------------------------------------------------------------
    # RECEIPT ITEMS
    #
    # GET  /api/receipts/<receipt_id>/items/
    # POST /api/receipts/<receipt_id>/items/
    # ------------------------------------------------------------------

    @action(
        detail=True,
        methods=["get", "post"],
        url_path="items",
    )
    def items(
        self,
        request,
        receipt_id=None,
    ):
        """
        List or add items belonging to a receipt.
        """

        receipt = self.get_object()

        # --------------------------------------------------------------
        # ADD ITEM
        # --------------------------------------------------------------

        if request.method == "POST":

            serializer = ReceiptItemSerializer(
                data=request.data
            )

            serializer.is_valid(
                raise_exception=True
            )

            # Receipt is taken from the URL.
            # Never trust a client-supplied receipt.
            item = serializer.save(
                receipt=receipt
            )

            # ReceiptItem.save() already recalculates
            # the receipt total.
            receipt.update_total()

            return Response(
                ReceiptItemSerializer(item).data,
                status=status.HTTP_201_CREATED,
            )

        # --------------------------------------------------------------
        # LIST ITEMS
        # --------------------------------------------------------------

        items = (
            receipt.items
            .select_related(
                "department",
                "subdepartment",
            )
            .all()
        )

        serializer = ReceiptItemSerializer(
            items,
            many=True,
        )

        return Response(
            serializer.data
        )

    # ------------------------------------------------------------------
    # EXPORT RECEIPT SUMMARY
    #
    # GET /api/receipts/export/
    # ------------------------------------------------------------------

    @action(
        detail=False,
        methods=["get"],
        url_path="export",
    )
    def export(
        self,
        request,
    ):
        """
        Export receipt summary to Excel.
        """

        receipts = self.filter_queryset(
            self.get_queryset()
        )

        date_from = request.query_params.get(
            "date_from"
        )

        date_to = request.query_params.get(
            "date_to"
        )

        if date_from:
            receipts = receipts.filter(
                date__gte=date_from
            )

        if date_to:
            receipts = receipts.filter(
                date__lte=date_to
            )

        headers = [
            "Receipt ID",
            "Date",
            "Store",
            "Payment Method",
            "Total Expense",
            "Recorded By",
        ]

        widths = [
            18,
            15,
            30,
            20,
            18,
            20,
        ]

        (
            workbook,
            worksheet,
            header_row,
        ) = _styled_workbook(
            "Receipt Summary",
            "Receipt Summary",
            headers,
            widths,
        )

        row = header_row + 1

        for receipt in receipts:

            worksheet.cell(
                row=row,
                column=1,
                value=receipt.receipt_id,
            )

            worksheet.cell(
                row=row,
                column=2,
                value=receipt.date,
            )

            worksheet.cell(
                row=row,
                column=3,
                value=receipt.store,
            )

            worksheet.cell(
                row=row,
                column=4,
                value=(
                    receipt
                    .get_payment_method_display()
                ),
            )

            total_cell = worksheet.cell(
                row=row,
                column=5,
                value=float(
                    receipt.total or 0
                ),
            )

            total_cell.number_format = (
                "#,##0.00"
            )

            worksheet.cell(
                row=row,
                column=6,
                value=receipt.recorded_by.username,
            )

            row += 1

        total = (
            receipts.aggregate(
                total=Sum("total")
            )["total"]
            or 0
        )

        worksheet.cell(
            row=row + 1,
            column=4,
            value="TOTAL",
        ).font = Font(
            bold=True
        )

        total_cell = worksheet.cell(
            row=row + 1,
            column=5,
            value=float(total),
        )

        total_cell.font = Font(
            bold=True
        )

        total_cell.number_format = (
            "#,##0.00"
        )

        response = HttpResponse(
            content_type=(
                "application/"
                "vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
        )

        response["Content-Disposition"] = (
            'attachment; '
            'filename="receipt_summary.xlsx"'
        )

        workbook.save(response)

        return response


# ===========================================================================
# RECEIPT ITEM VIEWSET
# ===========================================================================

class ReceiptItemViewSet(viewsets.ModelViewSet):
    """
    Full CRUD for ReceiptItem.

    Standard endpoints:

        GET    /api/receipt-items/
        POST   /api/receipt-items/
        GET    /api/receipt-items/<item_id>/
        PATCH  /api/receipt-items/<item_id>/
        PUT    /api/receipt-items/<item_id>/
        DELETE /api/receipt-items/<item_id>/

    Additional endpoints:

        GET /api/receipt-items/summary/
        GET /api/receipt-items/export/

    Filtering:

        /api/receipt-items/?receipt_id=RP-2026-00001

        /api/receipt-items/?department=5

        /api/receipt-items/?subdepartment=12

        /api/receipt-items/?department=5&subdepartment=12
    """

    queryset = (
        ReceiptItem.objects
        .select_related(
            "receipt",
            "receipt__recorded_by",
            "department",
            "subdepartment",
        )
        .all()
    )

    serializer_class = ReceiptItemSerializer

    permission_classes = [
        IsAuthenticated,
    ]

    lookup_url_kwarg = "item_id"

    # ------------------------------------------------------------------
    # QUERYSET / FILTERING
    # ------------------------------------------------------------------

    def get_queryset(self):

        queryset = super().get_queryset()

        receipt_id = (
            self.request.query_params.get(
                "receipt_id"
            )
        )

        date_from = (
            self.request.query_params.get(
                "date_from"
            )
        )

        date_to = (
            self.request.query_params.get(
                "date_to"
            )
        )

        department = (
            self.request.query_params.get(
                "department"
            )
        )

        subdepartment = (
            self.request.query_params.get(
                "subdepartment"
            )
        )

        if receipt_id:
            queryset = queryset.filter(
                receipt__receipt_id=receipt_id
            )

        if date_from:
            queryset = queryset.filter(
                receipt__date__gte=date_from
            )

        if date_to:
            queryset = queryset.filter(
                receipt__date__lte=date_to
            )

        if department:
            queryset = queryset.filter(
                department_id=department
            )

        if subdepartment:
            queryset = queryset.filter(
                subdepartment_id=subdepartment
            )

        return queryset

    # ------------------------------------------------------------------
    # CREATE
    # ------------------------------------------------------------------

    def create(
        self,
        request,
        *args,
        **kwargs,
    ):
        """
        Create an item for a specific receipt.

        React sends:

        {
            "receipt_id": "RP-2026-00001",
            "item_name": "...",
            "department": 5,
            "subdepartment": 12,
            "quantity": 2,
            "unit": "pieces",
            "unit_price": 500
        }

        The receipt is resolved server-side.
        """

        receipt_id = request.data.get(
            "receipt_id"
        )

        if not receipt_id:
            return Response(
                {
                    "receipt_id": [
                        "Receipt ID is required."
                    ]
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            receipt = Receipt.objects.get(
                receipt_id=receipt_id
            )

        except Receipt.DoesNotExist:
            return Response(
                {
                    "receipt_id": [
                        "Receipt does not exist."
                    ]
                },
                status=status.HTTP_404_NOT_FOUND,
            )

        # Copy the request data so we can safely remove
        # receipt_id before passing it to the serializer.
        data = request.data.copy()

        data.pop(
            "receipt_id",
            None,
        )

        serializer = self.get_serializer(
            data=data
        )

        serializer.is_valid(
            raise_exception=True
        )

        # Attach receipt from the server.
        item = serializer.save(
            receipt=receipt
        )

        # ReceiptItem.save() recalculates the total.
        receipt.update_total()

        headers = self.get_success_headers(
            serializer.data
        )

        return Response(
            ReceiptItemSerializer(item).data,
            status=status.HTTP_201_CREATED,
            headers=headers,
        )

    # ------------------------------------------------------------------
    # UPDATE
    # ------------------------------------------------------------------

    def perform_update(
        self,
        serializer,
    ):
        """
        ReceiptItem.save() automatically recalculates
        the parent receipt total after an update.
        """

        item = serializer.save()

        item.receipt.update_total()

    # ------------------------------------------------------------------
    # DELETE
    # ------------------------------------------------------------------

    def perform_destroy(
        self,
        instance,
    ):
        """
        ReceiptItem.delete() automatically recalculates
        the parent receipt total.
        """

        instance.delete()

    # ------------------------------------------------------------------
    # SUMMARY
    # ------------------------------------------------------------------

    @action(
        detail=False,
        methods=["get"],
        url_path="summary",
    )
    def summary(
        self,
        request,
    ):

        items = self.get_queryset()

        total_expenditure = (
            items.aggregate(
                total=Sum("total")
            )["total"]
            or 0
        )

        total_items = items.count()

        total_receipts = (
            items.values("receipt")
            .distinct()
            .count()
        )

        # --------------------------------------------------------------
        # DEPARTMENT SUMMARY
        # --------------------------------------------------------------

        department_summary = (
            items.values(
                "department",
                "department__name",
            )
            .annotate(
                total=Sum("total"),
                item_count=Count("id"),
                receipt_count=Count(
                    "receipt",
                    distinct=True,
                ),
            )
            .order_by("-total")
        )

        department_summary_data = []

        for row in department_summary:

            department_summary_data.append({
                "department": (
                    row["department__name"]
                ),
                "total": row["total"],
                "item_count": row["item_count"],
                "receipt_count": (
                    row["receipt_count"]
                ),
            })

        # --------------------------------------------------------------
        # ACTIVE DEPARTMENTS
        # --------------------------------------------------------------

        departments = (
            Department.objects
            .filter(is_active=True)
            .order_by("name")
        )

        return Response(
            {
                "total_expenditure": (
                    total_expenditure
                ),

                "total_items": total_items,

                "total_receipts": total_receipts,

                "department_summary": (
                    DepartmentSummarySerializer(
                        department_summary_data,
                        many=True,
                    ).data
                ),

                "date_from": (
                    request.query_params.get(
                        "date_from"
                    )
                ),

                "date_to": (
                    request.query_params.get(
                        "date_to"
                    )
                ),

                "selected_department": (
                    request.query_params.get(
                        "department"
                    )
                ),

                "selected_subdepartment": (
                    request.query_params.get(
                        "subdepartment"
                    )
                ),

                "departments": (
                    DepartmentSerializer(
                        departments,
                        many=True,
                    ).data
                ),
            }
        )

    # ------------------------------------------------------------------
    # EXCEL EXPORT
    # ------------------------------------------------------------------

    @action(
        detail=False,
        methods=["get"],
        url_path="export",
    )
    def export(
        self,
        request,
    ):

        items = self.get_queryset()

        headers = [
            "Receipt ID",
            "Date",
            "Store",
            "Item",
            "Department",
            "Subdepartment",
            "Quantity",
            "Unit",
            "Unit Price",
            "Total",
        ]

        widths = [
            18,
            15,
            30,
            30,
            25,
            25,
            12,
            15,
            18,
            18,
        ]

        (
            workbook,
            worksheet,
            header_row,
        ) = _styled_workbook(
            "Filtered Expenditure",
            "Filtered Expenditure Report",
            headers,
            widths,
        )

        row = header_row + 1

        for item in items:

            worksheet.cell(
                row=row,
                column=1,
                value=item.receipt.receipt_id,
            )

            worksheet.cell(
                row=row,
                column=2,
                value=item.receipt.date,
            )

            worksheet.cell(
                row=row,
                column=3,
                value=item.receipt.store,
            )

            worksheet.cell(
                row=row,
                column=4,
                value=item.item_name,
            )

            # Department is now a ForeignKey.
            worksheet.cell(
                row=row,
                column=5,
                value=item.department.name,
            )

            # Subdepartment is optional.
            worksheet.cell(
                row=row,
                column=6,
                value=(
                    item.subdepartment.name
                    if item.subdepartment
                    else ""
                ),
            )

            worksheet.cell(
                row=row,
                column=7,
                value=float(item.quantity),
            )

            worksheet.cell(
                row=row,
                column=8,
                value=item.get_unit_display(),
            )

            unit_price_cell = worksheet.cell(
                row=row,
                column=9,
                value=float(item.unit_price),
            )

            unit_price_cell.number_format = (
                "#,##0.00"
            )

            total_cell = worksheet.cell(
                row=row,
                column=10,
                value=float(item.total or 0),
            )

            total_cell.number_format = (
                "#,##0.00"
            )

            row += 1

        total = (
            items.aggregate(
                total=Sum("total")
            )["total"]
            or 0
        )

        worksheet.cell(
            row=row + 1,
            column=9,
            value="TOTAL",
        ).font = Font(
            bold=True
        )

        total_cell = worksheet.cell(
            row=row + 1,
            column=10,
            value=float(total),
        )

        total_cell.font = Font(
            bold=True
        )

        total_cell.number_format = (
            "#,##0.00"
        )

        response = HttpResponse(
            content_type=(
                "application/"
                "vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
        )

        response["Content-Disposition"] = (
            'attachment; '
            'filename="filtered_expenditure.xlsx"'
        )

        workbook.save(response)

        return response